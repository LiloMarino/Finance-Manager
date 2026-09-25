from __future__ import annotations

import warnings
from collections.abc import Iterator
from dataclasses import dataclass
from io import BytesIO

from openpyxl import load_workbook

# Extrato > Movimentação
MOVEMENT_COLUMNS = (
    "Entrada/Saída",
    "Data",
    "Movimentação",
    "Produto",
    "Quantidade",
    "Preço unitário",
    "Valor da Operação",
)

# Proventos > Recebidos
INCOME_COLUMNS = (
    "Produto",
    "Pagamento",
    "Tipo de Evento",
    "Quantidade",
    "Preço unitário",
    "Valor líquido",
)


@dataclass(frozen=True, slots=True, kw_only=True)
class B3Movement:
    """Uma linha de relatório da B3, com as células como texto. O relatório de
    proventos recebidos entra no mesmo formato: o pagamento é a data, o tipo de evento
    é a movimentação e o valor líquido é o valor."""

    direction: str
    movement_date: str
    movement: str
    product: str
    quantity: str
    unit_price: str
    amount: str


@dataclass(frozen=True, slots=True, kw_only=True)
class B3Report:
    """`income_only` é o relatório de proventos recebidos; o outro é o de
    movimentação."""

    income_only: bool
    movements: list[B3Movement]


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _rows(content: bytes) -> Iterator[list[str]]:
    """As linhas da primeira planilha, com as células como texto."""
    with warnings.catch_warnings():
        # O xlsx da B3 vem sem estilo padrão, e o openpyxl avisa a cada leitura
        warnings.simplefilter("ignore", UserWarning)
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    # O xlsx da B3 declara a dimensão A1:A1; sem recalcular, a leitura para na
    # primeira coluna
    sheet.reset_dimensions()
    for row in sheet.iter_rows(values_only=True):
        yield [_text(cell) for cell in row]


def read_b3_report(content: bytes) -> B3Report:
    """Lê o relatório pelo cabeçalho da linha 1, que diz qual dos dois ele é."""
    rows = _rows(content)
    header = next(rows, [])
    if all(column in header for column in MOVEMENT_COLUMNS):
        index = {column: header.index(column) for column in MOVEMENT_COLUMNS}
        return B3Report(
            income_only=False,
            movements=[
                B3Movement(
                    direction=cells[index["Entrada/Saída"]],
                    movement_date=cells[index["Data"]],
                    movement=cells[index["Movimentação"]],
                    product=cells[index["Produto"]],
                    quantity=cells[index["Quantidade"]],
                    unit_price=cells[index["Preço unitário"]],
                    amount=cells[index["Valor da Operação"]],
                )
                for cells in rows
                if any(cells)
            ],
        )
    if all(column in header for column in INCOME_COLUMNS):
        index = {column: header.index(column) for column in INCOME_COLUMNS}
        # O rodapé traz o total numa linha sem produto
        return B3Report(
            income_only=True,
            movements=[
                B3Movement(
                    direction="Credito",
                    movement_date=cells[index["Pagamento"]],
                    movement=cells[index["Tipo de Evento"]],
                    product=cells[index["Produto"]],
                    quantity=cells[index["Quantidade"]],
                    unit_price=cells[index["Preço unitário"]],
                    amount=cells[index["Valor líquido"]],
                )
                for cells in rows
                if cells[index["Produto"]]
            ],
        )
    raise ValueError(
        "Não é um relatório da B3: envie o de movimentação (Extrato) ou o de "
        "proventos recebidos."
    )
